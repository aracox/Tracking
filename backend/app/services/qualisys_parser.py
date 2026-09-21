"""Parser for Qualisys (QTM) tabular Excel exports.

Detected layout (identical for Position and Velocity exports):

    rows 0..N   metadata:  NO_OF_FRAMES, NO_OF_DATA_TYPES, FREQUENCY, TIME_STAMP,
                           DATA_INCLUDED, DATA_TYPES (list of column names)
    blank row
    header row  'Frame', 'Time', '<marker>_pos_X', '<marker>_pos_Y', '<marker>_pos_Z', ...
                (velocity files use '_vel_')
    data rows   integer frame number, time in seconds, then 3 columns per marker

The workbook has a single sheet. No unit fields are present. Values are read raw;
nothing is interpolated or modified here.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

_COL_RE = re.compile(r"^(?P<marker>.*)_(?P<kind>pos|vel)_(?P<axis>[XYZ])$")
_MAX_HEADER_SCAN = 60


class QualisysParseError(ValueError):
    pass


@dataclass
class RawTable:
    kind: str  # "pos" | "vel"
    metadata: dict[str, object]
    marker_names: list[str]  # display names (stripped)
    frames: np.ndarray  # [F] int64
    times: np.ndarray  # [F] float64
    values: np.ndarray  # [F, M, 3] float64, NaN for blank/non-numeric
    non_numeric_cells: int


def _to_float(x) -> float:
    if x is None or isinstance(x, bool):
        return np.nan
    if isinstance(x, (int, float)):
        return float(x)
    try:
        return float(str(x).strip().replace(",", "."))
    except ValueError:
        return np.nan


def _scan_header(rows_iter) -> tuple[dict[str, object], list, int]:
    """Consume rows until the 'Frame' header. Returns (metadata, header_row, rows_consumed)."""
    meta: dict[str, object] = {}
    for i, row in enumerate(rows_iter):
        if i >= _MAX_HEADER_SCAN:
            break
        first = row[0] if row else None
        if isinstance(first, str) and first.strip().lower() == "frame":
            return meta, list(row), i + 1
        if isinstance(first, str) and first.strip():
            vals = [c for c in row[1:] if c is not None]
            meta[first.strip()] = vals[0] if len(vals) == 1 else vals
    raise QualisysParseError("No 'Frame' header row found; not a Qualisys tabular export")


def _group_columns(header: list) -> tuple[str, list[str], list[str], list[tuple[int, int, int]]]:
    """Return (kind, raw_names, display_names, column-index triplets)."""
    kind: str | None = None
    order: list[str] = []
    cols: dict[str, dict[str, int]] = {}
    for idx, h in enumerate(header):
        if not isinstance(h, str):
            continue
        m = _COL_RE.match(h)
        if not m:
            continue
        kind = kind or m["kind"]
        if m["kind"] != kind:
            raise QualisysParseError("Mixed position/velocity columns in one file")
        name = m["marker"]
        if name not in cols:
            cols[name] = {}
            order.append(name)
        cols[name][m["axis"]] = idx
    if kind is None:
        raise QualisysParseError("No <marker>_pos_X / _vel_X columns found")
    incomplete = [n for n in order if set(cols[n]) != {"X", "Y", "Z"}]
    if incomplete:
        raise QualisysParseError(f"Markers without complete XYZ columns: {incomplete}")
    stripped = [n.strip() for n in order]
    display = stripped if len(set(stripped)) == len(stripped) else list(order)
    return kind, order, display, [(cols[n]["X"], cols[n]["Y"], cols[n]["Z"]) for n in order]


def read_header(path: Path) -> dict[str, object]:
    """Cheap header-only read for session listing (does not load data rows)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        meta, header, _ = _scan_header(ws.iter_rows(values_only=True))
        kind, _, display, _ = _group_columns(header)
        return {
            "kind": kind,
            "frames": int(meta.get("NO_OF_FRAMES", 0) or 0),
            "frequency": float(meta.get("FREQUENCY", 0) or 0),
            "markers": display,
        }
    finally:
        wb.close()


def parse_table(path: Path) -> RawTable:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        meta, header, _ = _scan_header(rows)
        kind, _, display, triplets = _group_columns(header)
        flat = np.array(triplets, dtype=np.int64).reshape(-1)
        frame_col = next(i for i, h in enumerate(header) if str(h).strip().lower() == "frame")
        time_col = next(
            (i for i, h in enumerate(header) if str(h).strip().lower() == "time"), None
        )
        frames: list[int] = []
        times: list[float] = []
        data: list[list[float]] = []
        bad = 0
        for row in rows:
            if not row or row[frame_col] is None:
                continue
            f = _to_float(row[frame_col])
            if np.isnan(f):
                continue
            frames.append(int(f))
            times.append(_to_float(row[time_col]) if time_col is not None else np.nan)
            vals = []
            for c in flat:
                cell = row[c] if c < len(row) else None
                v = _to_float(cell)
                if np.isnan(v) and cell is not None and str(cell).strip() != "":
                    bad += 1
                vals.append(v)
            data.append(vals)
    finally:
        wb.close()
    if not frames:
        raise QualisysParseError(f"{path.name}: no data rows")
    values = np.asarray(data, dtype=np.float64).reshape(len(frames), len(display), 3)
    return RawTable(
        kind=kind,
        metadata=meta,
        marker_names=display,
        frames=np.asarray(frames, dtype=np.int64),
        times=np.asarray(times, dtype=np.float64),
        values=values,
        non_numeric_cells=bad,
    )
